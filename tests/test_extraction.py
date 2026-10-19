"""Getting text out of PDF and DOCX files."""

import io

import pytest

from app.extraction.text import ResumeParseError, extract_text
from tests.fixtures.samples import SAMPLE_RESUMES


@pytest.mark.parametrize("sample", SAMPLE_RESUMES, ids=lambda s: s.filename)
def test_every_sample_resume_round_trips_through_its_real_file(sample, sample_resume_dir):
    with (sample_resume_dir / sample.filename).open("rb") as handle:
        text = extract_text(handle, sample.filename)

    # The name and a distinctive skill survive the PDF/DOCX round trip.
    assert sample.lines[0] in text
    assert sample.lines[-1].split(",")[0].strip() in text


def test_docx_table_contents_are_included(tmp_path):
    """Skills and education are often laid out in tables, not paragraphs."""
    import docx

    document = docx.Document()
    document.add_paragraph("Ananya Rao")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Education"
    table.rows[0].cells[1].text = "M.Tech in Computer Science"
    path = tmp_path / "tabular.docx"
    document.save(str(path))

    with path.open("rb") as handle:
        text = extract_text(handle, "tabular.docx")

    assert "M.Tech in Computer Science" in text


def test_ragged_pdf_whitespace_is_normalised(sample_resume_dir):
    with (sample_resume_dir / "aarav_mehta_backend.pdf").open("rb") as handle:
        text = extract_text(handle, "aarav_mehta_backend.pdf")

    assert "  " not in text  # no double spaces to break the duration regexes
    assert "\n\n\n" not in text


def test_unsupported_extension_is_rejected_by_name():
    with pytest.raises(ResumeParseError, match="Unsupported resume file type"):
        extract_text(io.BytesIO(b"plain text"), "resume.txt")


def test_corrupt_pdf_raises_rather_than_returning_empty_text():
    """The prototype swallowed this and returned '', so a broken upload looked
    identical to a blank resume."""
    with pytest.raises(ResumeParseError, match="Could not read the PDF"):
        extract_text(io.BytesIO(b"not really a pdf"), "broken.pdf")


def test_corrupt_docx_raises():
    with pytest.raises(ResumeParseError, match="Could not read the DOCX"):
        extract_text(io.BytesIO(b"PK\x03\x04 not really a docx"), "broken.docx")


def test_a_resume_with_no_extractable_text_says_so(tmp_path):
    """An image-only PDF reads as empty; the student needs to be told."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "scanned.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.showPage()
    pdf.save()

    with path.open("rb") as handle:
        with pytest.raises(ResumeParseError, match="No text could be extracted"):
            extract_text(handle, "scanned.pdf")
